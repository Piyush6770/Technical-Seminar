"""
Formative In-Place Table Column Width & Styling Adjuster for research_paper.docx
Preserves all user-made text, paragraphs, headings, and formatting changes.
Removes table section breaks so tables fit in the 2-column layout directly.
Sets exact column widths and font sizes for all 6 tables to fit within the 3.35 in column width.
"""

import os
import re
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCX_PATH = os.path.join(PROJECT_ROOT, "research_paper", "research_paper.docx")
DOCX_2COL_PATH = os.path.join(PROJECT_ROOT, "research_paper", "research_paper_ieee_2column.docx")

def set_cell_margins(table, top=20, bottom=20, left=30, right=30):
    tblPr = table._tbl.tblPr
    tblCellMar = tblPr.find(qn('w:tblCellMar'))
    if tblCellMar is None:
        tblCellMar = OxmlElement('w:tblCellMar')
        tblPr.append(tblCellMar)
    for m_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = tblCellMar.find(qn(f'w:{m_name}'))
        if node is None:
            node = OxmlElement(f'w:{m_name}')
            tblCellMar.append(node)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')

def set_table_width(table, width_dxa=4824):
    tblPr = table._tbl.tblPr
    tblW = tblPr.find(qn('w:tblW'))
    if tblW is None:
        tblW = OxmlElement('w:tblW')
        tblPr.append(tblW)
    tblW.set(qn('w:w'), str(width_dxa))
    tblW.set(qn('w:type'), 'dxa')

def format_tables_in_docx(doc_path):
    print(f"Processing: {doc_path}")
    doc = docx.Document(doc_path)
    body = doc._body._element

    # 1. Identify and remove empty sectPr paragraphs that were wrapping tables
    table_elements = [el for el in body if el.tag.endswith('tbl')]
    to_remove = []
    for tbl in table_elements:
        prev_el = tbl.getprevious()
        if prev_el is not None and prev_el.tag.endswith('p'):
            pPr = prev_el.find(qn('w:pPr'))
            if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
                runs_txt = ''.join([r.text for r in prev_el.findall('.//' + qn('w:t')) if r.text])
                if not runs_txt.strip():
                    to_remove.append(prev_el)
        
        curr = tbl.getnext()
        for _ in range(2):
            if curr is not None and curr.tag.endswith('p'):
                pPr = curr.find(qn('w:pPr'))
                if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
                    runs_txt = ''.join([r.text for r in curr.findall('.//' + qn('w:t')) if r.text])
                    if not runs_txt.strip():
                        if curr not in to_remove:
                            to_remove.append(curr)
                curr = curr.getnext()

    print(f"Removing {len(to_remove)} table-wrapping section break paragraphs...")
    for p in to_remove:
        body.remove(p)

    # 2. Define custom column widths (dxa) and font sizes for each table to fit 3.35 in (4824 dxa)
    # Total column width = 3.35 in = 4824 dxa
    TABLE_CONFIGS = [
        # Table 0: Forensic Paradigm Comparison (6 cols)
        {
            "widths_dxa": [1224, 648, 691, 720, 749, 792],
            "data_font": 6.0,
            "header_font": 6.4,
            "cell_pad": (15, 15, 25, 25)
        },
        # Table 1: 4-Stage Hierarchy (6 cols)
        {
            "widths_dxa": [605, 835, 432, 648, 1152, 1152],
            "data_font": 6.0,
            "header_font": 6.4,
            "cell_pad": (15, 15, 25, 25)
        },
        # Table 2: Dataset Splits (7 cols)
        {
            "widths_dxa": [864, 605, 547, 648, 648, 936, 576],
            "data_font": 5.8,
            "header_font": 6.2,
            "cell_pad": (12, 12, 20, 20)
        },
        # Table 3: Primary Benchmark Matrix E1-E6 (10 cols)
        {
            "widths_dxa": [403, 893, 389, 389, 403, 403, 389, 389, 576, 590],
            "data_font": 5.0,
            "header_font": 5.3,
            "cell_pad": (10, 10, 15, 15)
        },
        # Table 4: Baseline Comparison Across Stages (8 cols)
        {
            "widths_dxa": [720, 792, 547, 547, 547, 547, 547, 577],
            "data_font": 5.8,
            "header_font": 6.0,
            "cell_pad": (12, 12, 20, 20)
        },
        # Table 5: Appendix Hyperparameters (3 cols)
        {
            "widths_dxa": [1656, 1368, 1800],
            "data_font": 6.8,
            "header_font": 7.2,
            "cell_pad": (20, 20, 35, 35)
        }
    ]

    for t_idx, t in enumerate(doc.tables):
        if t_idx >= len(TABLE_CONFIGS):
            break
        cfg = TABLE_CONFIGS[t_idx]
        col_dxa = cfg["widths_dxa"]
        data_sz = cfg["data_font"]
        hdr_sz = cfg["header_font"]
        top_m, bot_m, l_m, r_m = cfg["cell_pad"]

        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        set_table_width(t, sum(col_dxa))
        set_cell_margins(t, top=top_m, bottom=bot_m, left=l_m, right=r_m)

        # Update tblGrid
        tblGrid = t._tbl.find(qn('w:tblGrid'))
        if tblGrid is not None:
            t._tbl.remove(tblGrid)
        tblGrid = OxmlElement('w:tblGrid')
        for w in col_dxa:
            gc = OxmlElement('w:gridCol')
            gc.set(qn('w:w'), str(w))
            tblGrid.append(gc)
        t._tbl.insert(t._tbl.index(t._tbl.tblPr) + 1, tblGrid)

        for r_idx, row in enumerate(t.rows):
            is_header = (r_idx == 0)
            target_font_sz = hdr_sz if is_header else data_sz
            for c_idx, cell in enumerate(row.cells):
                if c_idx < len(col_dxa):
                    cell.width = Inches(col_dxa[c_idx] / 1440.0)
                    tcPr = cell._tc.get_or_add_tcPr()
                    tcW = tcPr.find(qn('w:tcW'))
                    if tcW is None:
                        tcW = OxmlElement('w:tcW')
                        tcPr.append(tcW)
                    tcW.set(qn('w:w'), str(col_dxa[c_idx]))
                    tcW.set(qn('w:type'), 'dxa')

                # Adjust paragraph text formatting inside cell
                for p in cell.paragraphs:
                    p.paragraph_format.space_before = Pt(0.5)
                    p.paragraph_format.space_after = Pt(0.5)
                    p.paragraph_format.line_spacing = 1.0
                    for run in p.runs:
                        run.font.name = 'Times New Roman'
                        run.font.size = Pt(target_font_sz)
                        if is_header:
                            run.bold = True

    # 3. Resize all drawings (Figure 1) to column width (3.35 inches)
    target_cx = int(3.35 * 914400)
    for p in doc.paragraphs:
        drawings = p._p.findall('.//' + qn('w:drawing'))
        if drawings:
            extents = p._p.findall('.//' + qn('wp:extent'))
            for ext in extents:
                old_cx = int(ext.get('cx'))
                old_cy = int(ext.get('cy'))
                if old_cx > target_cx + 1000:
                    new_cx = target_cx
                    new_cy = int(old_cy * (new_cx / old_cx))
                    ext.set('cx', str(new_cx))
                    ext.set('cy', str(new_cy))
                    print(f"Resized diagram from {old_cx}x{old_cy} to {new_cx}x{new_cy}")
                    for a_ext in p._p.findall('.//' + qn('a:ext')):
                        if a_ext.get('cx') == str(old_cx):
                            a_ext.set('cx', str(new_cx))
                            a_ext.set('cy', str(new_cy))

    # 4. Remove any remaining empty sectPr paragraphs after the title/abstract header
    fig_sect_remove = []
    for el in body:
        if el.tag.endswith('p'):
            pPr = el.find(qn('w:pPr'))
            if pPr is not None and pPr.find(qn('w:sectPr')) is not None:
                runs_txt = ''.join([r.text for r in el.findall('.//' + qn('w:t')) if r.text])
                if not runs_txt.strip():
                    if body.index(el) > 4:
                        fig_sect_remove.append(el)
    
    for el in fig_sect_remove:
        body.remove(el)
    print(f"Removed {len(fig_sect_remove)} diagram sectPr paragraphs.")

    # 5. Save output
    try:
        doc.save(doc_path)
        print(f"Saved formatted document: {doc_path}")
    except PermissionError:
        alt_path = doc_path.replace('.docx', '_updated.docx')
        doc.save(alt_path)
        print(f"Warning: {doc_path} is currently locked by Word. Saved to: {alt_path}")

if __name__ == '__main__':
    format_tables_in_docx(DOCX_PATH)
    format_tables_in_docx(DOCX_2COL_PATH)
